from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document as DocxDocument


ROOT = Path(__file__).resolve().parents[1]
SAMPLES_DIR = ROOT / "samples"
CHINESE_FONT = Path("C:/Windows/Fonts/msyh.ttc")


def create_sample_docx() -> None:
    document = DocxDocument()
    document.add_heading("中文复杂文档解析示例", level=1)
    document.add_paragraph("本示例用于验证 DOCX 到 Markdown 和 JSON 的基础导出链路。")
    document.add_heading("项目目标", level=2)
    document.add_paragraph("保留标题层级、段落顺序、页码信息以及后续 RAG 所需的基础结构。")
    document.add_heading("适用场景", level=2)
    document.add_paragraph("适用于公文、报告、教辅材料等中文复杂文档的结构化处理。")
    table = document.add_table(rows=3, cols=3)
    cells = table.rows[0].cells
    cells[0].text = "模块"
    cells[1].text = "状态"
    cells[2].text = "说明"
    cells = table.rows[1].cells
    cells[0].text = "标题恢复"
    cells[1].text = "完成"
    cells[2].text = "支持 Heading 样式识别"
    cells = table.rows[2].cells
    cells[0].text = "表格抽取"
    cells[1].text = "完成"
    cells[2].text = "输出结构化 JSON"
    document.save(SAMPLES_DIR / "example.docx")


def create_sample_pdf() -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    font_name = "msyh"
    if CHINESE_FONT.exists():
        page.insert_font(fontname=font_name, fontfile=str(CHINESE_FONT))
    else:
        font_name = "helv"
    lines = [
        "中文复杂文档解析示例",
        "",
        "关于推进结构化解析能力建设的通知",
        "第 1 页",
        "一、项目背景",
        "本示例用于验证 PDF 到 Markdown 和 JSON 的基础导出链路。",
        "",
        "二、核心目标",
        "保留标题层级、段落顺序以及面向 RAG 的结构化内容。",
        "",
        "表 1 解析能力状态",
    ]
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontsize=12, fontname=font_name)
        y += 24
    x0, y0, width, height = 72, 320, 360, 120
    for row in range(4):
        y_line = y0 + row * (height / 3)
        page.draw_line((x0, y_line), (x0 + width, y_line))
    for col in range(4):
        x_line = x0 + col * (width / 3)
        page.draw_line((x_line, y0), (x_line, y0 + height))
    table_rows = [
        ["模块", "状态", "说明"],
        ["标题恢复", "完成", "支持版式启发式识别"],
        ["表格抽取", "完成", "输出结构化 JSON"],
    ]
    for row_index, row_values in enumerate(table_rows):
        for col_index, value in enumerate(row_values):
            page.insert_text(
                (x0 + col_index * (width / 3) + 8, y0 + row_index * (height / 3) + 24),
                value,
                fontsize=10,
                fontname=font_name,
            )
    pdf.save(SAMPLES_DIR / "example.pdf")
    pdf.close()


def main() -> None:
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    create_sample_docx()
    create_sample_pdf()


if __name__ == "__main__":
    main()
