from __future__ import annotations

import json
from pathlib import Path

from docx import Document as DocxDocument
from typer.testing import CliRunner

from zhdocparser.app import app


def test_cli_parse_exports_files_and_summary(tmp_path: Path) -> None:
    source = tmp_path / "demo.docx"
    output_dir = tmp_path / "outputs"
    doc = DocxDocument()
    doc.add_heading("项目说明", level=1)
    doc.add_paragraph("这是用于 CLI 测试的内容。")
    doc.save(source)

    runner = CliRunner()
    result = runner.invoke(app, ["parse", str(source), "-o", str(output_dir)])

    assert result.exit_code == 0
    assert "Summary:" in result.stdout

    bundle_dir = output_dir / "demo_docx"
    json_path = bundle_dir / "document.json"
    markdown_path = bundle_dir / "document.md"
    chunks_path = bundle_dir / "chunks.json"
    assert json_path.exists()
    assert markdown_path.exists()
    assert chunks_path.exists()

    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert payload["metadata"]["title"] == "项目说明"
    assert len(payload["chunks"]) == 1


def test_cli_parse_dir_exports_bundle_per_document(tmp_path: Path) -> None:
    source_dir = tmp_path / "sources"
    output_dir = tmp_path / "outputs"
    source_dir.mkdir()

    first_doc = DocxDocument()
    first_doc.add_heading("Notice", level=1)
    first_doc.add_paragraph("Alpha")
    first_doc.save(source_dir / "a.docx")

    second_doc = DocxDocument()
    second_doc.add_heading("Quarterly Report", level=1)
    second_doc.add_paragraph("Beta")
    second_doc.save(source_dir / "b.docx")

    runner = CliRunner()
    result = runner.invoke(app, ["parse-dir", str(source_dir), "-o", str(output_dir)])

    assert result.exit_code == 0
    assert "Batch Summary:" in result.stdout
    assert (output_dir / "a_docx" / "document.json").exists()
    assert (output_dir / "b_docx" / "chunks.json").exists()


def test_cli_parse_supports_doc_type_override_and_chunk_size(tmp_path: Path) -> None:
    source = tmp_path / "demo.docx"
    output_dir = tmp_path / "outputs"
    doc = DocxDocument()
    doc.add_heading("General Title", level=1)
    doc.add_paragraph("A" * 160)
    doc.add_paragraph("B" * 160)
    doc.save(source)

    runner = CliRunner()
    result = runner.invoke(
        app,
        [
            "parse",
            str(source),
            "-o",
            str(output_dir),
            "--doc-type",
            "report",
            "--max-chunk-chars",
            "150",
        ],
    )

    assert result.exit_code == 0
    payload = json.loads((output_dir / "demo_docx" / "document.json").read_text(encoding="utf-8"))
    chunks_payload = json.loads((output_dir / "demo_docx" / "chunks.json").read_text(encoding="utf-8"))
    assert payload["metadata"]["doc_type"] == "report"
    assert len(chunks_payload) >= 2
